"""Exportación a PDF (fpdf2) y DOCX (python-docx) desde markdown simple."""
from __future__ import annotations

import io
import re

from docx import Document
from docx.shared import Pt
from fpdf import FPDF


_PDF_REPLACEMENTS = {
    "—": "-", "–": "-", "•": "*", "→": "->", "←": "<-",
    "“": '"', "”": '"', "‘": "'", "’": "'", "…": "...",
    "€": "EUR", "©": "(c)", "®": "(R)", "™": "(TM)",
    "≤": "<=", "≥": ">=", "≠": "!=", "·": "-", "º": "o", "ª": "a",
}


def _strip_md_inline(line: str) -> str:
    line = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
    line = re.sub(r"\*(.+?)\*", r"\1", line)
    line = re.sub(r"`([^`]+)`", r"\1", line)
    return line


def _to_latin1(text: str) -> str:
    """Adapta texto para fpdf2 con fuente core (latin-1)."""
    for k, v in _PDF_REPLACEMENTS.items():
        text = text.replace(k, v)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def md_to_pdf(markdown_text: str, title: str = "Documento") -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(pdf.epw, 10, _to_latin1(title))
    pdf.ln(2)

    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()
        if not line:
            pdf.ln(3)
            continue
        if line.startswith("### "):
            pdf.set_font("Helvetica", "B", 12)
            pdf.multi_cell(pdf.epw, 7, _to_latin1(_strip_md_inline(line[4:])))
        elif line.startswith("## "):
            pdf.set_font("Helvetica", "B", 13)
            pdf.multi_cell(pdf.epw, 8, _to_latin1(_strip_md_inline(line[3:])))
        elif line.startswith("# "):
            pdf.set_font("Helvetica", "B", 14)
            pdf.multi_cell(pdf.epw, 9, _to_latin1(_strip_md_inline(line[2:])))
        elif line.lstrip().startswith(("- ", "* ")):
            pdf.set_font("Helvetica", "", 11)
            pdf.multi_cell(pdf.epw, 6, _to_latin1("* " + _strip_md_inline(line.lstrip()[2:])))
        else:
            pdf.set_font("Helvetica", "", 11)
            pdf.multi_cell(pdf.epw, 6, _to_latin1(_strip_md_inline(line)))

    out = pdf.output()
    if isinstance(out, str):
        return out.encode("latin-1", errors="replace")
    return bytes(out)


def md_to_docx(markdown_text: str, title: str = "Documento") -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading(title, level=0)

    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()
        if not line:
            doc.add_paragraph("")
            continue
        if line.startswith("### "):
            doc.add_heading(_strip_md_inline(line[4:]), level=3)
        elif line.startswith("## "):
            doc.add_heading(_strip_md_inline(line[3:]), level=2)
        elif line.startswith("# "):
            doc.add_heading(_strip_md_inline(line[2:]), level=1)
        elif line.lstrip().startswith(("- ", "* ")):
            doc.add_paragraph(_strip_md_inline(line.lstrip()[2:]), style="List Bullet")
        else:
            doc.add_paragraph(_strip_md_inline(line))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
